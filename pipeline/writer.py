"""Write the merged rows to a single-sheet .xlsx with platform-coloured column 1."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Config


class WriteError(Exception):
    """The output workbook could not be saved."""

SHEET_TITLE = "Entries"

HEADER_FILL = PatternFill("solid", fgColor="FF2F3E4E")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
THIN = Side(style="thin", color="FFD0D0D0")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MIN_WIDTH, MAX_WIDTH = 9, 34


def _column_widths(frame: pd.DataFrame) -> list[int]:
    widths = []
    for column in frame.columns:
        longest = max(
            [len(str(column))] + [len(str(v)) for v in frame[column].head(2000)] or [0]
        )
        widths.append(max(MIN_WIDTH, min(MAX_WIDTH, longest + 2)))
    return widths


def write_xlsx(frame: pd.DataFrame, path: Path | str, cfg: Config) -> Path:
    """Write `frame` (already ordered and deduplicated) to `path`."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    fills = {p.abbrev: PatternFill("solid", fgColor=p.fill_rgb) for p in cfg.platform_list}
    status_fills = {
        cfg.status.new_label: PatternFill("solid", fgColor=cfg.status.new_fill_rgb),
        cfg.status.existing_label: PatternFill(
            "solid", fgColor=cfg.status.existing_fill_rgb
        ),
    }
    status_index = (
        list(frame.columns).index(cfg.status.column) + 1
        if cfg.status.column in frame.columns
        else None
    )

    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_TITLE

    sheet.append(list(frame.columns))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER

    text_columns = {
        i + 1
        for i, column in enumerate(frame.columns)
        if column in {*cfg.id_columns, "Mobile Number", "DOB", "Payment Date"}
    }

    for values in frame.itertuples(index=False, name=None):
        sheet.append(["" if v is None else v for v in values])
        row = sheet.max_row
        for index in range(1, len(frame.columns) + 1):
            cell = sheet.cell(row=row, column=index)
            cell.border = CELL_BORDER
            if index in text_columns:
                cell.number_format = "@"
        platform_cell = sheet.cell(row=row, column=1)
        platform_cell.alignment = Alignment(horizontal="center")
        platform_cell.font = Font(bold=True)
        fill = fills.get(str(platform_cell.value))
        if fill is not None:
            platform_cell.fill = fill

        if status_index is not None:
            status_cell = sheet.cell(row=row, column=status_index)
            status_cell.alignment = Alignment(horizontal="center")
            status_fill = status_fills.get(str(status_cell.value))
            if status_fill is not None:
                status_cell.fill = status_fill
            if str(status_cell.value) == cfg.status.new_label:
                status_cell.font = Font(bold=True)

    for index, width in enumerate(_column_widths(frame), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    if len(frame):
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}"
        )

    try:
        book.save(path)
    except PermissionError as exc:
        raise WriteError(
            f"Could not write {path.name} -- it may be open in Excel, "
            "or the folder may be read-only. Close it and try again."
        ) from exc
    except OSError as exc:
        raise WriteError(f"Could not write {path}: {exc}") from exc
    return path
